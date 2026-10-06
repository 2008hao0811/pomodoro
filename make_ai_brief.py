# -*- coding: utf-8 -*-
"""生成 AI 每日简报 Word 文档（放到桌面）。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

import wecom_notify

OUT = r"C:\Users\Administrator\Desktop\AI每日简报_2026-10-06.docx"

# 新闻条目：每条 dict，含 title / points(list) / source / credibility
NEWS = [
    {
        "title": "OpenAI 洽谈 300 亿美元融资，阿联酋基金、贝莱德入局磋商",
        "points": [
            "据媒体报道，OpenAI 正洽谈新一轮约 300 亿美元融资。",
            "阿联酋主权基金与贝莱德（BlackRock）据报参与磋商。",
            "具体估值与是否最终落地以官方披露为准。",
        ],
        "source": "IT之家 https://www.ithome.com/1/009/961.htm",
        "credibility": "中（消息称 · 单一来源）",
    },
    {
        "title": "月之暗面（Kimi）完成上市前最后一轮融资：估值约 500 亿美元，拟赴港 IPO",
        "points": [
            "消息称月之暗面完成上市前最后一轮融资，估值约 500 亿美元。",
            "公司计划于明年一季度赴港交所 IPO。",
            "若成行，将成为国内大模型公司重要的资本化节点。",
        ],
        "source": "IT之家 https://www.ithome.com/1/009/971.htm",
        "credibility": "中（消息称 · 单一来源）",
    },
    {
        "title": "DeepSeek 接近完成至少 800 亿元融资，腾讯、宁德时代重金参与",
        "points": [
            "消息称 DeepSeek 接近完成至少 800 亿元新一轮融资。",
            "腾讯、宁德时代等据报重金参与。",
            "反映头部大模型公司在算力与商业化上加速布局。",
        ],
        "source": "IT之家 https://www.ithome.com/1/009/990.htm",
        "credibility": "中（消息称 · 单一来源）",
    },
    {
        "title": "智谱 GLM-5.3 上架亚马逊 AWS 大模型平台，打开海外收入分成通道",
        "points": [
            "智谱旗下大模型 GLM-5.3 已上架亚马逊 AWS 旗下大模型平台。",
            "此举为智谱打开海外市场的收入分成通道。",
            "国产大模型出海商业化迈出实质一步。",
        ],
        "source": "IT之家 https://www.ithome.com/1/009/946.htm",
        "credibility": "中（单一来源）",
    },
    {
        "title": "DeepSeek V4.1 Flash 发力，中美顶尖模型 LiveBench 跑分差距缩至 3%",
        "points": [
            "据 LiveBench 榜单，DeepSeek V4.1 Flash 表现突出。",
            "中美顶尖模型跑分差距缩小至约 3%。",
            "国产开源模型在通用能力上继续追赶前沿水平。",
        ],
        "source": "IT之家 https://www.ithome.com/1/009/751.htm",
        "credibility": "中（单一来源）",
    },
    {
        "title": "韩国拟推 4.7 万亿韩元专项计划，明年 3 月起研发前沿 AI 大模型",
        "points": [
            "韩国政府拟推出 4.7 万亿韩元（约合人民币 250 亿元级）专项计划。",
            "计划自明年 3 月起研发前沿 AI 大模型。",
            "各国在主权 AI 与基础模型能力上的竞争持续升温。",
        ],
        "source": "IT之家 https://www.ithome.com/1/010/000.htm",
        "credibility": "中（单一来源）",
    },
]


def set_cn_font(run, name="微软雅黑", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, size=11, bold=False, color=None, align=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_cn_font(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def main():
    doc = Document()

    # 标题
    add_para(doc, "AI 每日简报", size=22, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "2026 年 10 月 6 日 · 星期二", size=12, color=(0x88, 0x88, 0x88),
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    # 图例
    add_para(doc, "可信度图例：高 = 多来源交叉验证　中 = 单一权威来源　低 = 存疑/待确认",
             size=9, color=(0x99, 0x99, 0x99), space_after=12)

    if not NEWS:
        add_para(doc, "（简报内容待检索结果填入——当前新闻条目为空。）",
                 size=11, color=(0xcc, 0x00, 0x00))
        doc.save(OUT)
        print("SKELETON SAVED:", OUT)
        return

    for i, item in enumerate(NEWS, 1):
        title = item.get("title", "")
        points = item.get("points", [])
        source = item.get("source", "")
        cred = item.get("credibility", "中")

        add_para(doc, f"{i}. {title}", size=13, bold=True, space_after=2)
        for p in points:
            add_para(doc, "　• " + p, size=11, space_after=2)
        if source:
            add_para(doc, "来源：" + source, size=9, color=(0x33, 0x66, 0xcc), space_after=1)
        add_para(doc, "可信度：" + cred, size=9, color=(0x66, 0x66, 0x66), space_after=10)

    add_para(doc, "说明：以上条目整理自 IT之家今日（2026-10-06）科技资讯，均为单一来源、未做多源交叉验证；"
                  "标注「消息称」的条目属媒体报道/传闻，请以官方公告为准。",
             size=9, color=(0x99, 0x99, 0x99), space_after=0)

    doc.save(OUT)
    print("SAVED:", OUT)
    wecom_notify.notify_brief_generated("AI 每日简报", OUT)


if __name__ == "__main__":
    main()
