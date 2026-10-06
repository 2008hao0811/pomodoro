# -*- coding: utf-8 -*-
"""生成全球金融每日简报 Word 文档（放到桌面）。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

import wecom_notify

OUT = r"C:\Users\Administrator\Desktop\全球金融每日简报_2026-10-06.docx"

NEWS = [
    {
        "title": "美国长债收益率飙至 2002 年来新高，30 年期逼近 6%「债务陷阱」警戒线",
        "points": [
            "10 年期报 5.34%、30 年期报 5.7%，均创 2002 年以来最高。",
            "ISM 服务业价格指数跳升至 74（2022 年 7 月以来最高）；市场对美联储 10 月加息定价约 25%、12 月已完全计入一次 25bp 加息。",
            "BMO 固收主管判断 30 年期本月内破 6%「不可避免」，届时美国借贷成本或超名义 GDP 增速。",
        ],
        "source": "华尔街见闻 wallstreetcn.com/articles/3783051 · 新浪财经 finance.sina.com.cn",
        "credibility": "高（两独立来源，数据一致 · 事实）",
    },
    {
        "title": "日本央行行长植田和男暗示准备继续加息",
        "points": [
            "称潜在通胀存在超调至 2% 目标上方的风险（受能源价格、AI 需求、日元疲软驱动）。",
            "强调「金融环境依然宽松」，将继续提高借贷成本。",
        ],
        "source": "新浪财经 finance.sina.com.cn（转引外媒快讯）",
        "credibility": "中（单一来源 · 事实，转引）",
    },
    {
        "title": "美股三大指数集体收涨，纳指创收盘历史新高逼近 27500 点",
        "points": [
            "道指 +0.18%、标普 500 +0.66%、纳指 +1.05%（收 27477.31）。",
            "英伟达、台积电涨超 2% 再创新高，SpaceX 涨超 7%；中概股普涨，纳斯达克中国金龙指数 +1.71%。",
        ],
        "source": "东方财富 finance.eastmoney.com · 新浪财经",
        "credibility": "高（两独立来源，收盘点位一致 · 事实）",
    },
    {
        "title": "欧元跌至 17 个月新低，法国债务状况威胁欧元区稳定",
        "points": [
            "欧元/美元跌破 1.12（触及 1.1161）。",
            "勒庞提出赤字削减计划并警告法国或丧失「金融自主权」，市场担忧法国成「欧债风暴中心」。",
        ],
        "source": "东方财富 · 华尔街见闻 wallstreetcn.com/articles/3783053",
        "credibility": "高（两来源一致 · 汇率事实 / 定性为分析）",
    },
    {
        "title": "美国战略石油储备跌至 1982 年最低 + 特朗普放宽免税柴油，油价承压回落",
        "points": [
            "SPR 降至 2.83 亿桶，创 1982 年以来最低。",
            "特朗普签署行政令临时允许红柴油上高速并递延燃油税；WTI 收跌约 2% 至 89.3 美元/桶，布伦特失守 100 美元。",
        ],
        "source": "CNBC cnbc.com · 新浪财经 · 东方财富",
        "credibility": "高（CNBC 原文 + 两个中文来源互证 · 事实）",
    },
    {
        "title": "OpenAI 推进约 300 亿美元融资，阿联酋 MGX 牵头财团",
        "points": [
            "阿联酋多支基金拟合计出资最高 100 亿美元，贝莱德参与，并与 Thrive、a16z 磋商。",
            "交易未完成，细节仍可能变动。",
        ],
        "source": "新浪财经（转引彭博等）",
        "credibility": "中（单一来源 · 待确认）",
    },
    {
        "title": "快手旗下可灵 AI 筹备港股 IPO，拟募资至少 10 亿美元",
        "points": [
            "与中金、高盛、瑞银合作，最快明年（2027 年初递表）上市。",
            "公司估值已突破 1200 亿元人民币。",
        ],
        "source": "新浪财经 · 东方财富 finance.eastmoney.com",
        "credibility": "高（两独立来源一致 · 据知情人士）",
    },
    {
        "title": "（补充）美国国会敦促美联储审查香港金管局 FIMA 美元回购便利准入",
        "points": [
            "美国众议院「对华竞争特设委员会」主席致信美联储，CNBC 独家获得信件。",
            "反映美国对华金融施压新方向；美联储已收到信件并将回复。",
        ],
        "source": "CNBC cnbc.com（独家）",
        "credibility": "中（CNBC 独家单一来源 · 事实）",
    },
]


def set_cn_font(run, name="微软雅黑", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, size=11, bold=False, color=None, align=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_cn_font(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def main():
    doc = Document()

    add_para(doc, "全球金融每日简报", size=22, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_para(doc, "2026 年 10 月 6 日 · 星期二", size=12, color=(0x88, 0x88, 0x88),
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    add_para(doc, "可信度图例：高 = 多来源交叉验证　中 = 单一权威来源　低 = 存疑/待确认",
             size=9, color=(0x99, 0x99, 0x99), space_after=12)

    for i, item in enumerate(NEWS, 1):
        title = item.get("title", "")
        points = item.get("points", [])
        source = item.get("source", "")
        cred = item.get("credibility", "中")

        add_para(doc, f"{i}. {title}", size=13, bold=True, space_after=2)
        for p in points:
            add_para(doc, "　• " + p, size=11, space_after=2)
        if source:
            add_para(doc, "来源：" + source, size=9, color=(0x33, 0x66, 0xcc), space_after=1)
        add_para(doc, "可信度：" + cred, size=9, color=(0x66, 0x66, 0x66), space_after=10)

    add_para(doc, "说明：美股/商品数据为美东时间 10 月 5 日（周一）收盘，在亚太时区 10 月 6 日（今日）报道；"
                  "10/6 美股盘中走势本简报未覆盖。Reuters/Bloomberg/FT/WSJ 直接抓取超时，"
                  "国际新闻多经中文权威财经媒体转引，跨独立外媒互证有限，已按实标注可信度。",
             size=9, color=(0x99, 0x99, 0x99), space_after=0)

    doc.save(OUT)
    print("SAVED:", OUT)
    wecom_notify.notify_brief_generated("全球金融每日简报", OUT)


if __name__ == "__main__":
    main()
